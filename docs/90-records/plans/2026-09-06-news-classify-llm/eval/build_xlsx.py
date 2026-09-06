"""Gộp 3 file nháp Opus + bài mẫu thành gold-draft.xlsx cho chủ dự án rà. Chạy: uv run --with openpyxl python build_xlsx.py"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

G = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(G, "gold-draft-2026-09-06.xlsx")
SUBS = {"1a": "Thể chế và văn bản pháp quy", "1b": "Điều hành Chính phủ", "1c": "Tiền tệ và tỷ giá", "1d": "Đầu tư công và hạ tầng", "1e": "Số liệu vĩ mô",
        "1f": "Thuế và ngân sách", "2a": "Chứng khoán thế giới", "2b": "Ngân hàng trung ương", "2c": "Hàng hoá và năng lượng", "2d": "An ninh và địa chính trị",
        "2e": "Thương mại và thuế quan", "3a": "CBTT và sự kiện quyền", "3b": "Giao dịch nội bộ và cổ đông lớn", "3c": "Vốn và cấu trúc", "3d": "KQKD và vận hành",
        "3e": "Nhận định và diễn biến thị trường", "3f": "Phái sinh, chứng quyền, ETF/quỹ", "3g": "Vi phạm và xử phạt", "3h": "Margin và ký quỹ",
        "3i": "Xếp hạng tín nhiệm và ESG", "x": "Loại bỏ (xã hội, thể thao, PR…)"}
inds = json.load(open(os.path.join(G, "industries.json"), encoding="utf-8"))
articles, drafts = {}, {}
for i in (1, 2, 3):
    for l in open(os.path.join(G, f"chunk-{i}.jsonl"), encoding="utf-8"):
        r = json.loads(l); articles[r["article_id"]] = r
    for l in open(os.path.join(G, f"draft-{i}.jsonl"), encoding="utf-8"):
        r = json.loads(l); drafts[r["article_id"]] = r
missing = [a for a in articles if a not in drafts]
assert not missing, f"thiếu nháp: {missing}"

wb = Workbook()
ws = wb.active; ws.title = "gold"
head = ["stt", "article_id", "nguồn", "feed", "nhóm gợi ý", "ngày đăng", "tiêu đề", "sapo", "đoạn đầu", "url",
        "NHÁP nhóm", "NHÁP sub", "NHÁP mã", "NHÁP ngành", "khó?", "ghi chú nháp",
        "CHỐT nhóm", "CHỐT sub", "CHỐT mã", "CHỐT ngành", "ghi chú của chủ dự án"]
ws.append(head)
hard_fill = PatternFill("solid", fgColor="FFF2CC"); final_fill = PatternFill("solid", fgColor="E2EFDA")
for col in range(1, len(head) + 1):
    c = ws.cell(row=1, column=col); c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    if col >= 17: c.fill = final_fill
order = [json.loads(l)["article_id"] for i in (1, 2, 3) for l in open(os.path.join(G, f"chunk-{i}.jsonl"), encoding="utf-8")]
for n, aid in enumerate(order, 1):
    a, d = articles[aid], drafts[aid]
    tk, ind = " ".join(d["tickers"]), " ".join(d["industries"])
    ws.append([n, aid, a["source"], a["feed"], a["hint"] if a["hint"] is not None else "", (a["published_at"] or "")[:10], a["title"], a["sapo"] or "",
               (a["excerpt"] or "")[:700], a["url"], d["group"], d["sub"], tk, ind, "x" if d["hard"] else "", d["note"],
               d["group"], d["sub"], tk, ind, ""])
    r = ws.max_row
    for col in range(1, len(head) + 1):
        c = ws.cell(row=r, column=col); c.alignment = Alignment(wrap_text=True, vertical="top")
        if col >= 17: c.fill = final_fill
    if d["hard"]:
        for col in (11, 12, 13, 14, 15, 16): ws.cell(row=r, column=col).fill = hard_fill
widths = [5, 9, 11, 14, 7, 11, 40, 40, 60, 30, 8, 8, 18, 26, 6, 30, 8, 8, 18, 26, 30]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "H2"
last = ws.max_row
dv_g = DataValidation(type="list", formula1='"1,2,3,x"', allow_blank=False); dv_g.add(f"Q2:Q{last}"); ws.add_data_validation(dv_g)
dv_s = DataValidation(type="list", formula1=f'"{",".join(SUBS)}"', allow_blank=False); dv_s.add(f"R2:R{last}"); ws.add_data_validation(dv_s)

w2 = wb.create_sheet("ma_sub"); w2.append(["sub", "tên"]); [w2.append([k, v]) for k, v in SUBS.items()]
w3 = wb.create_sheet("ma_nganh"); w3.append(["code", "tên"]); [w3.append([i["code"], i["name"]]) for i in inds]
for w in (w2, w3): w.column_dimensions["A"].width = 14; w.column_dimensions["B"].width = 45
w4 = wb.create_sheet("huong_dan")
for line in [
    "Bộ đánh giá (gold set) lưới phân loại tin — bản NHÁP do Opus gán, 150 bài lấy ngẫu nhiên 20/08–06/09/2026 (38/38/38/36 theo nhóm gợi ý feed).",
    "Chỉ sửa 5 cột xanh (CHỐT nhóm · CHỐT sub · CHỐT mã · CHỐT ngành · ghi chú). Các cột khác là ngữ cảnh, đừng sửa.",
    "Cột CHỐT đã điền sẵn = NHÁP; chỉ sửa chỗ anh thấy sai. Bài tô vàng ở phần NHÁP là bài Opus tự nhận khó — nên đọc kỹ hơn.",
    "nhóm: 1 = chủ thể Việt Nam · 2 = chủ thể nước ngoài/thế giới · 3 = doanh nghiệp niêm yết / TTCK Việt Nam · x = không phải tin kinh tế. Khi nhóm = x thì sub = x.",
    "sub: xem sheet ma_sub. mã: các mã niêm yết cách nhau bằng dấu cách, tối đa 5, quan trọng nhất trước, chỉ khi nhóm 3; để trống nếu không có.",
    "ngành: mã trong sheet ma_nganh, cách nhau bằng dấu cách, tối đa 3, áp cho mọi nhóm; để trống nếu bài vĩ mô thuần hoặc x.",
    "Lưu lại đúng file .xlsx này (không đổi tên cột). Trợ lý sẽ đọc 5 cột CHỐT để tạo gold.jsonl và chấm model.",
]: w4.append([line])
w4.column_dimensions["A"].width = 160
wb.save(OUT)
print("saved", OUT, "rows", last - 1, "hard", sum(1 for d in drafts.values() if d["hard"]))
