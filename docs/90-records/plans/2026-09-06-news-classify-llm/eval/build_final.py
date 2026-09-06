"""Gộp p1/p2/adj → gold.jsonl (đáp án chuẩn) + gold-2026-09-06.xlsx (chủ dự án soát) + thống kê. Chạy: uv run --with openpyxl python build_final.py OUTDIR"""
import json, os, sys, collections
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
G = os.path.dirname(os.path.abspath(__file__)); OUT = sys.argv[1]
def loadp(prefix, n):
    d = {}
    for i in range(1, n + 1):
        p = f"{G}/{prefix}-{i}.jsonl"
        if os.path.exists(p):
            for l in open(p, encoding="utf-8"): r = json.loads(l); d[r["article_id"]] = r
    return d
art = {}; order = []
for i in range(1, 9):
    for l in open(f"{G}/chunk-{i}.jsonl", encoding="utf-8"): r = json.loads(l); art[r["article_id"]] = r; order.append(r["article_id"])
p1, p2, adj = loadp("p1", 8), loadp("p2", 8), loadp("adj", 3)
inds = {i["code"]: i["name"] for i in json.load(open(f"{G}/../industries.json", encoding="utf-8"))}
final, src = {}, collections.Counter()
for aid in order:
    a, b = p1[aid], p2[aid]
    same = a["group"] == b["group"] and a["sub"] == b["sub"] and set(a["tickers"]) == set(b["tickers"]) and set(a["industries"]) == set(b["industries"])
    if same: final[aid] = {k: a[k] for k in ("group", "sub", "tickers", "industries")}; final[aid]["how"] = "đồng thuận"; src["đồng thuận"] += 1
    else:
        assert aid in adj, f"thiếu phân xử {aid}"; x = adj[aid]
        final[aid] = {k: x[k] for k in ("group", "sub", "tickers", "industries")}; final[aid]["how"] = f"phân xử→{x['picked']}"; src[f"phân xử→{x['picked']}"] += 1
# Luật nhất quán (c) chốt 2026-09-06 sau phân xử: tin lãi suất/tỷ giá/ngân hàng trung ương (1c, 2b) luôn gắn NGANHANG (chủ dự án: Fed ⇒ NGANHANG đúng)
fixed_c = 0
for aid, f_ in final.items():
    if f_["sub"] in ("1c", "2b") and "NGANHANG" not in f_["industries"]:
        f_["industries"] = (["NGANHANG"] + f_["industries"])[:3]; fixed_c += 1
src["luật (c) NGANHANG áp thêm"] = fixed_c
with open(f"{OUT}/gold.jsonl", "w", encoding="utf-8") as f:
    for aid in order: f.write(json.dumps({"article_id": aid, **{k: final[aid][k] for k in ("group", "sub", "tickers", "industries")}}, ensure_ascii=False) + "\n")
wb = Workbook(); ws = wb.active; ws.title = "gold"
head = ["stt", "article_id", "nguồn", "nhóm gợi ý", "tiêu đề", "sapo", "đoạn đầu", "url", "L1 nhóm/sub", "L1 mã", "L1 ngành", "L1 ghi chú", "L2 nhóm/sub", "L2 mã", "L2 ngành", "L2 ghi chú",
        "cách chốt", "lý do phân xử", "CHỐT nhóm", "CHỐT sub", "CHỐT mã", "CHỐT ngành", "ghi chú chủ dự án"]
ws.append(head); ff = PatternFill("solid", fgColor="E2EFDA"); hf = PatternFill("solid", fgColor="FFF2CC")
for c in range(1, len(head) + 1): ws.cell(row=1, column=c).font = Font(bold=True)
for n, aid in enumerate(order, 1):
    a, b, x, fl = p1[aid], p2[aid], adj.get(aid), final[aid]; r = art[aid]
    ws.append([n, aid, r["source"], r["hint"] if r["hint"] is not None else "", r["title"], r["sapo"] or "", (r["excerpt"] or "")[:600], r["url"],
               f'{a["group"]}/{a["sub"]}', " ".join(a["tickers"]), " ".join(a["industries"]), a["note"], f'{b["group"]}/{b["sub"]}', " ".join(b["tickers"]), " ".join(b["industries"]), b["note"],
               fl["how"], x["reason"] if x else "", fl["group"], fl["sub"], " ".join(fl["tickers"]), " ".join(fl["industries"]), ""])
    rr = ws.max_row
    for c in range(1, len(head) + 1):
        ws.cell(row=rr, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        if c >= 19: ws.cell(row=rr, column=c).fill = ff
    if x:
        for c in range(9, 19): ws.cell(row=rr, column=c).fill = hf
for i, w in enumerate([5, 9, 11, 7, 40, 38, 50, 28, 9, 16, 22, 26, 9, 16, 22, 26, 13, 34, 8, 8, 16, 22, 26], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "F2"
w2 = wb.create_sheet("thong_ke"); ag = json.load(open(f"{G}/agreement.json")); n = ag["n"]
for k, v in ag["agree"].items(): w2.append([f"đồng thuận {k}", f"{v}/{n} = {v/n:.0%}"])
for k, v in src.items(): w2.append([k, v])
w2.append(["phân bố nhóm chốt", json.dumps(collections.Counter(f["group"] for f in final.values()))])
w2.append(["phân bố sub chốt", json.dumps(collections.Counter(f["sub"] for f in final.values()), ensure_ascii=False)])
w2.append(["phân bố ngành chốt", json.dumps(collections.Counter(c for f in final.values() for c in f["industries"]), ensure_ascii=False)])
w2.column_dimensions["A"].width = 24; w2.column_dimensions["B"].width = 140
w3 = wb.create_sheet("huong_dan")
for line in ["Bộ gold 400 bài: hai lượt Opus gán độc lập, lượt ba phân xử chỗ lệch. Dòng tô vàng = hai lượt lệch nhau, đã phân xử (đọc 'lý do phân xử').",
             "gold.jsonl là đáp án dùng để chấm. Muốn sửa: sửa 4 cột xanh CHỐT rồi báo trợ lý dựng lại gold.jsonl.",
             "Taxonomy bản 2026-09-06: 2f Doanh nghiệp và kinh tế các nước; 3e = thị trường tài sản trong nước (chứng khoán, vàng, BĐS); x xét nội dung; phân nhóm theo chủ thể."]: w3.append([line])
w3.column_dimensions["A"].width = 160
wb.save(f"{OUT}/gold-2026-09-06.xlsx")
print("gold", len(final), dict(src), "nhóm", collections.Counter(f["group"] for f in final.values()))
